from __future__ import annotations
import io, re
from pathlib import Path
import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(page_title='Proteasome QC & Replicate Reviewer', page_icon='🔬', layout='wide')
st.title('🔬 Proteasome QC & Technical-Replicate Reviewer')
st.caption('Review mass calls manually, compare technical replicates, and export a curated QC workbook.')

QC_OPTIONS = ['Not reviewed','Keep / confident','Interesting / investigate','Low intensity','Noise','Ambiguous','Exclude']


def parse_peak(s):
    m=re.search(r'Peak\s*(\d+)', str(s), re.I); return f'Peak {int(m.group(1))}' if m else ''

def infer_meta(name, report=''):
    s=f'{name} {report}'.lower()
    age='Old/Adult' if any(x in s for x in ['old','adult','aged']) else ('Infant/Young' if any(x in s for x in ['infant','young','juvenile','pup']) else '')
    organ=''
    for k,v in {'liver':'Liver','brain':'Brain','lung':'Lung','kidney':'Kidney','kideny':'Kidney','heart':'Heart','spleen':'Spleen'}.items():
        if k in s: organ=v; break
    return age,organ

def parse_combined(uploaded):
    raw=uploaded.getvalue(); xls=pd.ExcelFile(io.BytesIO(raw), engine='openpyxl')
    if 'Combined results' not in xls.sheet_names: raise ValueError('Workbook has no "Combined results" sheet.')
    grid=pd.read_excel(io.BytesIO(raw), sheet_name='Combined results', header=None, engine='openpyxl')
    rows=[]; report=''; datafile=''; peak=''; headers=None
    for i in range(len(grid)):
        vals=grid.iloc[i].tolist(); first=str(vals[0]) if pd.notna(vals[0]) else ''
        if '|  Peak' in first or re.search(r'\.docx.*Peak\s*\d+', first, re.I):
            report=first.split('|')[0].strip(); peak=parse_peak(first); headers=None; continue
        if first=='Report file':
            report=str(vals[1]) if len(vals)>1 and pd.notna(vals[1]) else report; peak=peak or parse_peak(report); continue
        if first=='Data file':
            datafile=str(vals[1]) if len(vals)>1 and pd.notna(vals[1]) else ''; continue
        if first=='Subunit' and 'Detected_Mass_Da' in vals:
            headers=[str(v) if pd.notna(v) else '' for v in vals]; continue
        if headers and 'Detected_Mass_Da' in headers:
            d=dict(zip(headers,vals)); mass=d.get('Detected_Mass_Da')
            if pd.notna(mass) and isinstance(mass,(int,float,np.number)):
                age,organ=infer_meta(uploaded.name,report)
                d.update({'Source_Workbook':uploaded.name,'Report_File':report,'Data_File_Name':datafile,'Peak':peak or parse_peak(report),'Age_Group':age,'Organ':organ})
                rows.append(d)
    if not rows: raise ValueError('No detected masses were found in Combined results.')
    out=pd.DataFrame(rows)
    for c in ['Detected_Mass_Da','Calculated_Mass_Da','Difference_Da','Difference_ppm','Charge_State_Count','Sum_Intensity','Relative_Intensity_pct']:
        if c in out: out[c]=pd.to_numeric(out[c],errors='coerce')
    return out

def default_review(df):
    x=df.copy(); x['Manual_QC']='Not reviewed'; x['Reviewer_Notes']=''; return x

def greedy_match(a,b,tol,same_peak=False):
    A=a.reset_index(drop=True); B=b.reset_index(drop=True); pairs=[]; used=set()
    order=A.sort_values('Detected_Mass_Da').index
    for ia in order:
        ma=float(A.loc[ia,'Detected_Mass_Da']); cand=[]
        for ib,row in B.iterrows():
            if ib in used: continue
            if same_peak and str(A.loc[ia,'Peak'])!=str(row.get('Peak','')): continue
            d=abs(ma-float(row['Detected_Mass_Da']))
            if d<=tol: cand.append((d,ib))
        if cand:
            d,ib=min(cand); used.add(ib); pairs.append((ia,ib,d))
        else: pairs.append((ia,None,np.nan))
    for ib in B.index:
        if ib not in used: pairs.append((None,ib,np.nan))
    rec=[]
    for ia,ib,d in pairs:
        ra=A.loc[ia] if ia is not None else None; rb=B.loc[ib] if ib is not None else None
        m1=float(ra['Detected_Mass_Da']) if ra is not None else np.nan; m2=float(rb['Detected_Mass_Da']) if rb is not None else np.nan
        status='Reproduced' if ia is not None and ib is not None else ('Replicate 1 only' if ia is not None else 'Replicate 2 only')
        rec.append({'Consensus_Mass_Da':np.nanmean([m1,m2]),'Rep1_Mass_Da':m1,'Rep2_Mass_Da':m2,'Abs_Difference_Da':d,'Replicate_Status':status,
                    'Rep1_Peak':ra.get('Peak','') if ra is not None else '','Rep2_Peak':rb.get('Peak','') if rb is not None else '',
                    'Rep1_Category':ra.get('Category','') if ra is not None else '','Rep2_Category':rb.get('Category','') if rb is not None else '',
                    'Rep1_Subunit':ra.get('Subunit','') if ra is not None else '','Rep2_Subunit':rb.get('Subunit','') if rb is not None else '',
                    'Rep1_Modification':ra.get('Modification','') if ra is not None else '','Rep2_Modification':rb.get('Modification','') if rb is not None else '',
                    'Rep1_Sum_Intensity':ra.get('Sum_Intensity',np.nan) if ra is not None else np.nan,'Rep2_Sum_Intensity':rb.get('Sum_Intensity',np.nan) if rb is not None else np.nan,
                    'Rep1_Charge_State_Count':ra.get('Charge_State_Count',np.nan) if ra is not None else np.nan,'Rep2_Charge_State_Count':rb.get('Charge_State_Count',np.nan) if rb is not None else np.nan,
                    'Rep1_Manual_QC':ra.get('Manual_QC','') if ra is not None else '','Rep2_Manual_QC':rb.get('Manual_QC','') if rb is not None else ''})
    return pd.DataFrame(rec)

def style_ws(ws):
    ws.sheet_view.showGridLines=False; ws.freeze_panes='A2'
    fill=PatternFill('solid',fgColor='D9EAF7')
    for c in ws[1]: c.font=Font(bold=True); c.fill=fill; c.alignment=Alignment(wrap_text=True)
    for col in range(1,ws.max_column+1):
        letter=get_column_letter(col); vals=[len(str(ws.cell(r,col).value)) for r in range(1,min(ws.max_row,300)+1) if ws.cell(r,col).value is not None]
        ws.column_dimensions[letter].width=min(max(vals+[10])+2,38)

def export_xlsx(reviewed,comparison,qualified,recurring):
    bio=io.BytesIO()
    with pd.ExcelWriter(bio,engine='openpyxl') as w:
        reviewed.to_excel(w,index=False,sheet_name='Reviewed masses')
        comparison.to_excel(w,index=False,sheet_name='Replicate comparison')
        qualified.to_excel(w,index=False,sheet_name='Qualified masses')
        recurring.to_excel(w,index=False,sheet_name='Recurring candidates')
        for ws in w.book.worksheets: style_ws(ws)
    return bio.getvalue()

st.sidebar.header('Replicate comparison')
tol=st.sidebar.number_input('Mass tolerance (± Da)',min_value=0.05,value=1.0,step=0.05)
same_peak=st.sidebar.checkbox('Require same LC peak',value=False)
st.sidebar.caption('A replicate-only mass is flagged for review; it is not automatically declared noise.')

st.subheader('1. Upload analysis workbooks')
files=st.file_uploader('Upload one workbook for manual QC, or two workbooks for technical-replicate comparison',type=['xlsx'],accept_multiple_files=True)
if not files: st.stop()
if len(files)>2: st.warning('This version compares the first two workbooks as a technical-replicate pair. All uploaded files can be supported in a future multi-replicate mode.')
parsed=[]
for f in files[:2]:
    try: parsed.append((f.name,parse_combined(f)))
    except Exception as e: st.error(f'{f.name}: {e}')
if not parsed: st.stop()

st.subheader('2. Manual RAW-data QC')
reviewed_parts=[]
for idx,(name,df) in enumerate(parsed):
    st.markdown(f'### {name}')
    review=default_review(df)
    showcols=[c for c in ['Age_Group','Organ','Peak','Detected_Mass_Da','Category','Subunit','Accession','Modification','Possible_Explanation','Charge_State_Range','Charge_State_Count','Sum_Intensity','Relative_Intensity_pct','Manual_QC','Reviewer_Notes','Report_File','Data_File_Name'] if c in review.columns]
    edited=st.data_editor(review[showcols],hide_index=True,use_container_width=True,num_rows='fixed',key=f'editor_{idx}',
        column_config={'Manual_QC':st.column_config.SelectboxColumn('Manual QC',options=QC_OPTIONS,required=True),'Reviewer_Notes':st.column_config.TextColumn('Reviewer notes',width='large')})
    # merge edited manual fields back using row order
    full=review.copy(); full['Manual_QC']=edited['Manual_QC'].values; full['Reviewer_Notes']=edited['Reviewer_Notes'].values
    reviewed_parts.append(full)

reviewed=pd.concat(reviewed_parts,ignore_index=True)
comparison=pd.DataFrame()
if len(reviewed_parts)>=2:
    st.subheader('3. Technical-replicate comparison')
    comparison=greedy_match(reviewed_parts[0],reviewed_parts[1],tol,same_peak)
    st.dataframe(comparison,use_container_width=True,hide_index=True)
    c1,c2,c3=st.columns(3)
    c1.metric('Reproduced',int((comparison.Replicate_Status=='Reproduced').sum()))
    c2.metric('Replicate 1 only',int((comparison.Replicate_Status=='Replicate 1 only').sum()))
    c3.metric('Replicate 2 only',int((comparison.Replicate_Status=='Replicate 2 only').sum()))
else:
    st.info('Upload a second workbook from the technical repeat to activate replicate comparison.')

# Qualification: manual Noise/Exclude/Low intensity are removed; others retained. Replicate-only remains reviewable, not auto-noise.
excluded={'Noise','Exclude','Low intensity'}
qualified=reviewed[~reviewed['Manual_QC'].isin(excluded)].copy()
# Recurring candidates = reproduced pairs where either app category is non-Matched; manual noise/exclude blocks them.
if not comparison.empty:
    recurring=comparison[comparison['Replicate_Status'].eq('Reproduced')].copy()
    nonmatch=~(recurring['Rep1_Category'].eq('Matched') & recurring['Rep2_Category'].eq('Matched'))
    good=~recurring['Rep1_Manual_QC'].isin(excluded) & ~recurring['Rep2_Manual_QC'].isin(excluded)
    recurring=recurring[nonmatch & good].reset_index(drop=True)
else:
    recurring=pd.DataFrame(columns=['Consensus_Mass_Da','Replicate_Status'])

st.subheader('4. Curated outputs')
a,b,c=st.columns(3)
a.metric('Reviewed observations',len(reviewed)); b.metric('Qualified after manual QC',len(qualified)); c.metric('Recurring non-matched candidates',len(recurring))
with st.expander('Qualified masses'): st.dataframe(qualified,use_container_width=True,hide_index=True)
with st.expander('Recurring candidates',expanded=True): st.dataframe(recurring,use_container_width=True,hide_index=True)

st.subheader('5. Download reviewed workbook')
out=export_xlsx(reviewed,comparison,qualified,recurring)
st.download_button('Download QC + replicate analysis Excel',data=out,file_name='proteasome_QC_replicate_review.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
