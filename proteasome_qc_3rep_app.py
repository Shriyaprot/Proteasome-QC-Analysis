from __future__ import annotations

import io
import re
from pathlib import Path

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Proteasome Replicate QC", page_icon="🔬", layout="wide")
st.title("🔬 Proteasome Technical-Replicate QC")
st.caption("Compare 2 or 3 technical-replicate analysis workbooks and append only Recurring candidates and Non-recurring candidates to a copy of replicate 1.")

QC_OPTIONS = [
    "Not reviewed", "Keep / confident", "Interesting / investigate",
    "Low intensity", "Noise", "Ambiguous", "Exclude",
]


def parse_peak(value):
    m = re.search(r"Peak\s*(\d+)", str(value), re.I)
    return f"Peak {int(m.group(1))}" if m else ""


def infer_meta(filename, report=""):
    text = f"{filename} {report}".lower()
    age = ""
    if any(x in text for x in ["old", "adult", "aged"]):
        age = "Old/Adult"
    elif any(x in text for x in ["infant", "young", "juvenile", "pup"]):
        age = "Infant/Young"
    organ = ""
    for key, label in {
        "liver": "Liver", "brain": "Brain", "lung": "Lung", "kidney": "Kidney",
        "kideny": "Kidney", "heart": "Heart", "spleen": "Spleen",
    }.items():
        if key in text:
            organ = label
            break
    return age, organ


def parse_combined(uploaded, rep_no):
    raw = uploaded.getvalue()
    xls = pd.ExcelFile(io.BytesIO(raw), engine="openpyxl")
    if "Combined results" not in xls.sheet_names:
        raise ValueError(f'{uploaded.name}: no "Combined results" sheet.')
    grid = pd.read_excel(io.BytesIO(raw), sheet_name="Combined results", header=None, engine="openpyxl")
    rows, report, datafile, peak, headers = [], "", "", "", None
    for i in range(len(grid)):
        vals = grid.iloc[i].tolist()
        first = str(vals[0]).strip() if pd.notna(vals[0]) else ""
        if "|  Peak" in first or re.search(r"\.docx.*Peak\s*\d+", first, re.I):
            report = first.split("|")[0].strip()
            peak = parse_peak(first)
            headers = None
            continue
        if first == "Report file":
            report = str(vals[1]) if len(vals) > 1 and pd.notna(vals[1]) else report
            peak = peak or parse_peak(report)
            continue
        if first == "Data file":
            datafile = str(vals[1]) if len(vals) > 1 and pd.notna(vals[1]) else ""
            continue
        if first == "Subunit" and "Detected_Mass_Da" in vals:
            headers = [str(v) if pd.notna(v) else "" for v in vals]
            continue
        if headers and "Detected_Mass_Da" in headers:
            d = dict(zip(headers, vals))
            mass = d.get("Detected_Mass_Da")
            if pd.notna(mass) and isinstance(mass, (int, float, np.number)):
                age, organ = infer_meta(uploaded.name, report)
                d.update({
                    "Replicate": rep_no, "Source_Workbook": uploaded.name,
                    "Report_File": report, "Data_File_Name": datafile,
                    "Peak": peak or parse_peak(report), "Age_Group": age, "Organ": organ,
                })
                rows.append(d)
    if not rows:
        raise ValueError(f"{uploaded.name}: no detected masses found in Combined results.")
    out = pd.DataFrame(rows)
    for col in ["Detected_Mass_Da", "Calculated_Mass_Da", "Difference_Da", "Difference_ppm",
                "Charge_State_Count", "Sum_Intensity", "Relative_Intensity_pct"]:
        if col in out.columns:
            out[col] = pd.to_numeric(out[col], errors="coerce")
    return out


def best_assignment(row):
    sub = row.get("Subunit", "")
    mod = row.get("Modification", "")
    cat = row.get("Category", "")
    expl = row.get("Possible_Explanation", "")
    parts = [str(x).strip() for x in [sub, mod] if pd.notna(x) and str(x).strip()]
    assignment = " — ".join(parts)
    if not assignment and pd.notna(expl) and str(expl).strip():
        assignment = str(expl).strip()
    return assignment, cat


def same_peak_ok(cluster, row, require_same_peak):
    if not require_same_peak:
        return True
    peaks = [str(x.get("Peak", "")) for x in cluster["members"] if str(x.get("Peak", ""))]
    rp = str(row.get("Peak", ""))
    return not peaks or not rp or all(p == rp for p in peaks)


def cluster_replicates(frames, tolerance, require_same_peak=False):
    """Greedily form one-observation-per-replicate mass clusters within tolerance."""
    clusters = []
    # Seed with replicate 1.
    for _, row in frames[0].iterrows():
        clusters.append({"members": [row.to_dict()]})

    # Add each later replicate to nearest eligible existing cluster; unmatched rows seed new clusters.
    for rep_idx, frame in enumerate(frames[1:], start=2):
        candidates = []
        for ri, row in frame.iterrows():
            mass = float(row["Detected_Mass_Da"])
            for ci, cl in enumerate(clusters):
                if any(int(m["Replicate"]) == rep_idx for m in cl["members"]):
                    continue
                if not same_peak_ok(cl, row, require_same_peak):
                    continue
                center = float(np.mean([float(m["Detected_Mass_Da"]) for m in cl["members"]]))
                diff = abs(mass - center)
                if diff <= tolerance:
                    candidates.append((diff, ri, ci))
        candidates.sort(key=lambda x: x[0])
        used_rows, used_clusters = set(), set()
        for diff, ri, ci in candidates:
            if ri in used_rows or ci in used_clusters:
                continue
            clusters[ci]["members"].append(frame.loc[ri].to_dict())
            used_rows.add(ri)
            used_clusters.add(ci)
        for ri, row in frame.iterrows():
            if ri not in used_rows:
                clusters.append({"members": [row.to_dict()]})
    return clusters


def cluster_to_row(cluster, n_reps):
    members = sorted(cluster["members"], key=lambda x: int(x["Replicate"]))
    masses = np.array([float(m["Detected_Mass_Da"]) for m in members], dtype=float)
    out = {
        "Consensus_Mass_Da": float(np.mean(masses)),
        "Mass_SD_Da": float(np.std(masses, ddof=1)) if len(masses) >= 2 else np.nan,
        "Mass_Range_Da": float(np.max(masses) - np.min(masses)) if len(masses) >= 2 else np.nan,
        "Detected_In_n": len(members),
        "Total_Replicates": n_reps,
        "Reproducibility": f"{len(members)}/{n_reps}",
    }
    by_rep = {int(m["Replicate"]): m for m in members}
    for rep in range(1, n_reps + 1):
        m = by_rep.get(rep)
        prefix = f"Rep{rep}_"
        if m is None:
            for field in ["Mass_Da", "App_Category", "Assignment", "Peak", "Sum_Intensity",
                          "Relative_Intensity_pct", "Charge_State_Range", "Charge_State_Count", "Report_File"]:
                out[prefix + field] = np.nan if field in ["Mass_Da", "Sum_Intensity", "Relative_Intensity_pct", "Charge_State_Count"] else ""
            continue
        assignment, category = best_assignment(m)
        out.update({
            prefix + "Mass_Da": m.get("Detected_Mass_Da", np.nan),
            prefix + "App_Category": category,
            prefix + "Assignment": assignment,
            prefix + "Peak": m.get("Peak", ""),
            prefix + "Sum_Intensity": m.get("Sum_Intensity", np.nan),
            prefix + "Relative_Intensity_pct": m.get("Relative_Intensity_pct", np.nan),
            prefix + "Charge_State_Range": m.get("Charge_State_Range", ""),
            prefix + "Charge_State_Count": m.get("Charge_State_Count", np.nan),
            prefix + "Report_File": m.get("Report_File", ""),
        })
    first = members[0]
    out["Age_Group"] = first.get("Age_Group", "")
    out["Organ"] = first.get("Organ", "")
    out["Manual_QC"] = "Not reviewed"
    out["Reviewer_Notes"] = ""
    return out


def build_tables(frames, tolerance, same_peak, min_reps):
    clusters = cluster_replicates(frames, tolerance, same_peak)
    rows = [cluster_to_row(c, len(frames)) for c in clusters]
    all_df = pd.DataFrame(rows).sort_values("Consensus_Mass_Da").reset_index(drop=True)
    recurring = all_df[all_df["Detected_In_n"] >= min_reps].copy()
    nonrecurring = all_df[all_df["Detected_In_n"] < min_reps].copy()
    return recurring, nonrecurring


def style_sheet(ws):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        vals = [len(str(ws.cell(r, col).value)) for r in range(1, min(ws.max_row, 400) + 1) if ws.cell(r, col).value is not None]
        ws.column_dimensions[letter].width = min(max(vals + [10]) + 2, 38)


def append_dataframe_sheet(wb, name, df):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    for c_idx, col in enumerate(df.columns, 1):
        ws.cell(1, c_idx, col)
    for r_idx, values in enumerate(df.itertuples(index=False, name=None), 2):
        for c_idx, value in enumerate(values, 1):
            if pd.isna(value):
                value = None
            ws.cell(r_idx, c_idx, value)
    style_sheet(ws)


def export_into_rep1(rep1_bytes, recurring, nonrecurring):
    wb = load_workbook(io.BytesIO(rep1_bytes))
    append_dataframe_sheet(wb, "Recurring candidates", recurring)
    append_dataframe_sheet(wb, "Non-recurring candidates", nonrecurring)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


st.sidebar.header("Comparison settings")
tolerance = st.sidebar.number_input("Mass tolerance (± Da)", min_value=0.05, value=1.0, step=0.05)
same_peak = st.sidebar.checkbox("Require same LC peak", value=False)
st.sidebar.caption("Masses are grouped across replicates using the selected tolerance. Replicate-only does not automatically mean noise.")

st.subheader("1. Upload 2 or 3 technical-replicate analysis workbooks")
files = st.file_uploader(
    "Select replicate Excel files in order (Rep1, Rep2, Rep3)",
    type=["xlsx"], accept_multiple_files=True,
)
if len(files) < 2:
    st.info("Upload at least 2 technical replicates. You can upload up to 3.")
    st.stop()
if len(files) > 3:
    st.error("Please upload a maximum of 3 technical replicates in this version.")
    st.stop()

min_default = 2 if len(files) >= 2 else 1
min_reps = st.sidebar.number_input(
    "Minimum replicates required to call recurring",
    min_value=2, max_value=len(files), value=min_default, step=1,
)

try:
    frames = [parse_combined(f, i + 1) for i, f in enumerate(files)]
except Exception as exc:
    st.error(str(exc))
    st.stop()

recurring, nonrecurring = build_tables(frames, tolerance, same_peak, int(min_reps))

st.subheader("2. Recurring candidates")
st.caption(f"Detected in at least {int(min_reps)}/{len(files)} technical replicates within ±{tolerance:g} Da. Mass SD is calculated only from your replicate measurements.")
if recurring.empty:
    recurring_edit = recurring.copy()
    st.info("No recurring masses were found with the current settings.")
else:
    recurring_edit = st.data_editor(
        recurring, hide_index=True, use_container_width=True, num_rows="fixed", key="recurring_editor",
        column_config={
            "Manual_QC": st.column_config.SelectboxColumn("Manual QC", options=QC_OPTIONS, required=True),
            "Reviewer_Notes": st.column_config.TextColumn("Reviewer notes", width="large"),
        },
    )

st.subheader("3. Non-recurring candidates")
st.caption(f"Detected in fewer than {int(min_reps)} of {len(files)} replicates. These are review candidates, not automatically noise.")
if nonrecurring.empty:
    nonrecurring_edit = nonrecurring.copy()
    st.info("All mass clusters satisfy the current recurrence requirement.")
else:
    nonrecurring_edit = st.data_editor(
        nonrecurring, hide_index=True, use_container_width=True, num_rows="fixed", key="nonrecurring_editor",
        column_config={
            "Manual_QC": st.column_config.SelectboxColumn("Manual QC", options=QC_OPTIONS, required=True),
            "Reviewer_Notes": st.column_config.TextColumn("Reviewer notes", width="large"),
        },
    )

st.subheader("4. Summary")
cols = st.columns(len(frames) + 2)
for i, frame in enumerate(frames):
    cols[i].metric(f"Replicate {i+1} masses", len(frame))
cols[-2].metric("Recurring clusters", len(recurring_edit))
cols[-1].metric("Non-recurring clusters", len(nonrecurring_edit))

st.subheader("5. Download")
st.write("The download is a copy of **replicate 1** with only two QC sheets appended/replaced: **Recurring candidates** and **Non-recurring candidates**. All original analysis sheets are preserved.")
output = export_into_rep1(files[0].getvalue(), recurring_edit, nonrecurring_edit)
base = re.sub(r"\.xlsx$", "", files[0].name, flags=re.I)
st.download_button(
    "Download replicate 1 + QC sheets", data=output,
    file_name=f"{base}_with_replicate_QC.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
